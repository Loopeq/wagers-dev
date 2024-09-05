from openpyxl.worksheet.worksheet import Worksheet
from src.xml.utils import period_corrector
from components import draw_title, draw_row, fill_cell, get_refs
from src.xml.styles import H2, H3
from openpyxl.styles import Alignment, PatternFill
from src.xml.charter.chart_components import draw_chart_head, draw_chart, draw_chart_head_points
from src.xml.components import draw_title


def _create_without_points(ws, crds, dates, local_data, type_keys):
    head_crds, home_crds, away_crds = draw_chart_head(ws, crds, dates, local_data)
    home_values_y = get_refs(ws, crds=(home_crds[0], home_crds[1] - 1), width=len(local_data[type_keys[0]]),
                             height=0)

    away_values_y = get_refs(ws, crds=(away_crds[0], away_crds[1] - 1), width=len(local_data[type_keys[1]]),
                             height=0)

    values_x = get_refs(ws, crds=(head_crds[0] + 1, head_crds[1] - 1), width=len(dates), height=0)

    chart, anchor = draw_chart(data_y=[home_values_y], data_x=values_x, crds=away_crds,
                               colors=['008000'], pos_letter='A')
    ws.add_chart(chart, anchor)
    chart, anchor = draw_chart(data_y=[away_values_y], data_x=values_x, crds=away_crds,
                               colors=['B22222'], pos_letter='K')
    ws.add_chart(chart, anchor)

    chart, anchor = draw_chart(data_y=[home_values_y, away_values_y], data_x=values_x, crds=away_crds,
                               colors=['008000', 'B22222'], pos_letter='U', with_labels=False)
    ws.add_chart(chart, anchor)


def _create_with_points(ws, crds, dates, local_data, type_keys):
    h_point_keys = list(local_data[type_keys[0]].keys())
    a_point_keys = list(local_data[type_keys[1]].keys())
    crds = (crds[0], crds[1])
    for index, _ in enumerate(h_point_keys):
        h_point = h_point_keys[index]
        a_point = a_point_keys[index]

        ref_data = {type_keys[0]: local_data[type_keys[0]][h_point], type_keys[1]: local_data[type_keys[1]][a_point]}
        head_crds, home_crds, away_crds = draw_chart_head_points(ws, crds, dates, ref_data, points=(h_point, a_point))
        home_values_y = get_refs(ws, crds=(home_crds[0], home_crds[1] - 1), width=len(local_data[type_keys[0]]),
                                 height=0)

        away_values_y = get_refs(ws, crds=(away_crds[0], away_crds[1] - 1), width=len(local_data[type_keys[1]]),
                                 height=0)

        values_x = get_refs(ws, crds=(head_crds[0] + 1, head_crds[1] - 1), width=len(dates), height=0)

        chart, anchor = draw_chart(data_y=[home_values_y], data_x=values_x, crds=away_crds,
                                   colors=['008000'], pos_letter='A')
        ws.add_chart(chart, anchor)
        chart, anchor = draw_chart(data_y=[away_values_y], data_x=values_x, crds=away_crds,
                                   colors=['B22222'], pos_letter='J')
        ws.add_chart(chart, anchor)
        chart, anchor = draw_chart(data_y=[home_values_y, away_values_y], data_x=values_x, crds=away_crds,
                                   colors=['008000', 'B22222'], pos_letter='T', with_labels=False)
        ws.add_chart(chart, anchor)
        crds = (crds[0], crds[1] + 25)
    return crds


def create_line_chart(ws: Worksheet, dates: list[str], data: dict, title: str,
                      crds: tuple):
    period_keys = data.keys()

    total_crds = (crds[0], crds[1])
    dates = list(map(lambda m: m.split()[1], dates))
    for pk in period_keys:
        cur_title = period_corrector(title, pk)
        local_data = data[pk]
        type_keys = list(data[pk].keys())
        crds = draw_title(ws, total_crds, cur_title, width=20, font=H3, height=20,
                          alignment=Alignment(horizontal='center', vertical='center'),
                          background_color='ffff00')
        if isinstance(local_data[type_keys[0]], list):
            _create_without_points(ws, crds, dates, local_data, type_keys)
            total_crds = (total_crds[0], total_crds[1] + 28)
        else:
            point_crds = _create_with_points(ws, crds, dates, local_data, type_keys)
            total_crds = (point_crds[0], total_crds[1] + point_crds[1])
